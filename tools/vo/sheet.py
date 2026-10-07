#!/usr/bin/env python3
"""The ElevenLabs guide as a spreadsheet: one row per file to generate, one tab per day.

Needs openpyxl.  Upload the .xlsx to Google Drive and open it with Google Sheets.

    python3 tools/vo/sheet.py
"""
import os, sys
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from vo import DAYS, VOICES, BREAK, BRANCH_NAMES, blocks_for

OUT = os.path.join('docs', 'scripts', 'IELTS Speaking - ElevenLabs audio.xlsx')
TABS = {'prototype': 'Day 1', 'prototype-day2': 'Day 2', 'prototype-day3': 'Day 3', 'prototype-mock': 'Mock test'}
TINT = {'sarah': 'FCEFD6', 'nolan': 'DCE6F2', 'aarav': 'E3F1E0', 'student': 'F6E1EC'}

HEAD = Font(bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor='2B2B2B')
WRAP = Alignment(wrap_text=True, vertical='top')
TOP = Alignment(vertical='top')
LINE = Border(bottom=Side(style='thin', color='DDDDDD'))

def header(ws, cols):
    ws.append([c for c, _ in cols])
    for k, (_, w) in enumerate(cols, 1):
        cell = ws.cell(row=1, column=k)
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, Alignment(vertical='center')
        ws.column_dimensions[cell.column_letter].width = w
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = 'A2'

def main():
    wb = Workbook()
    how = wb.active; how.title = 'How to record'
    days = [(d, t, *blocks_for(d)) for d, t in DAYS]
    total = sum(len(b) for *_, b in days)

    rows = [
        ['IELTS Speaking · Audio for ElevenLabs'],
        [f'{total} files to generate, one per script section. Each day has its own tab.'],
        [],
        ['How to record'],
        ['1', 'In ElevenLabs Text to Speech, choose the model Eleven Multilingual v2. It reads the <break time="1.5s" /> tags as pauses.'],
        ['2', 'Pick the voice in the Voice column. Use the same settings for every file of that voice, on every day. Start with Stability 50%, Similarity 75%, Style 0%, Speaker boost on, speed 1.0.'],
        ['3', 'Click the cell in "Text to paste", copy it, and paste it into ElevenLabs exactly as it is: same words, same punctuation, break tags included. Generate.'],
        ['4', 'Listen. If any line is wrong or a pause is missing, regenerate the whole thing. Don\'t edit the audio.'],
        ['5', 'Download the MP3 and rename it to the name in "Save as".'],
        ['6', 'Put it in that day\'s folder (below), or send the files to Claude. Tick the Done column as you go.'],
        [],
        ['Why the pauses matter', 'Each file gets cut at its pauses into one clip per line. Each clip\'s length decides when that line\'s caption and animation happen, so keep the pauses and nothing drifts out of step.'],
        [],
        ['Voice', 'Direction'],
    ]
    rows += [[name, note] for v, (name, note) in VOICES.items()]
    rows += [[], ['Tab', 'Files', 'Save into']]
    rows += [[TABS[d], len(b), f'{d}/vo/raw/'] for d, t, _, b in days]
    for r in rows: how.append(r)
    how.column_dimensions['A'].width = 22
    how.column_dimensions['B'].width = 110
    how.column_dimensions['C'].width = 26
    for r in how.iter_rows():
        for c in r: c.alignment = WRAP
    how['A1'].font = Font(bold=True, size=16)
    for ref in ('A4', 'A12', 'A14', 'B14', 'A20', 'B20', 'C20'):
        how[ref].font = Font(bold=True)
    for k, v in enumerate(VOICES, 15):
        how.cell(row=k, column=1).fill = PatternFill('solid', fgColor=TINT[v])

    cols = [('Done', 7), ('#', 5), ('Section', 26), ('Theme', 12), ('Voice', 10),
            ('Save as', 20), ('Lines', 7), ('Text to paste', 90)]
    for d, t, data, blocks in days:
        ws = wb.create_sheet(TABS[d])
        header(ws, cols)
        for b in blocks:
            text = f' {BREAK} '.join(l['say'] for l in b['lines'])
            ws.append(['', f"{b['beat']:02d}", b['title'], BRANCH_NAMES.get(b['branch'], '') if b['branch'] else '',
                       VOICES[b['voice']][0], b['file'], len(b['lines']), text])
            r = ws.max_row
            for k in range(1, len(cols) + 1):
                c = ws.cell(row=r, column=k)
                c.alignment = WRAP if k in (3, 8) else TOP
                c.border = LINE
            ws.cell(row=r, column=5).fill = PatternFill('solid', fgColor=TINT[b['voice']])
            ws.cell(row=r, column=6).font = Font(name='Courier New')

    # every clip, for checking what a file gets cut into
    ws = wb.create_sheet('Clips (reference)')
    header(ws, [('Tab', 10), ('#', 5), ('Section', 26), ('Theme', 12), ('Voice', 10),
                ('From file', 20), ('Line', 5), ('Words', 80), ('Clip', 22)])
    for d, t, data, blocks in days:
        for b in blocks:
            for k, l in enumerate(b['lines'], 1):
                ws.append([TABS[d], f"{b['beat']:02d}", b['title'],
                           BRANCH_NAMES.get(b['branch'], '') if b['branch'] else '',
                           VOICES[b['voice']][0], b['file'], k, l['say'], l['files'][0]])
                for c in ws[ws.max_row]: c.alignment = WRAP

    wb.save(OUT)
    print(OUT, total, 'files')

if __name__ == '__main__':
    main()
